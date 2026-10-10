"""Create a synthetic fixture and audit the standard cover / TOC / image cells."""
import sys
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

target = Path(sys.argv[2])
if sys.argv[1] == 'create':
    target.parent.mkdir(parents=True, exist_ok=True)
    picture = target.parent / 'diagram.png'
    image = Image.new('RGB', (1000, 330), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((45, 95, 260, 245), outline='#19487d', width=5)
    draw.rectangle((720, 95, 955, 245), outline='#19487d', width=5)
    draw.line((260, 170, 720, 170), fill='#19487d', width=6)
    draw.text((370, 130), '20 m / 4 m / 10.000 m', fill='black')
    image.save(picture)
    doc = Document()
    doc.sections[0].page_width, doc.sections[0].page_height = Cm(21), Cm(29.7)
    for style in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        doc.styles[style].font.name = 'Microsoft YaHei'
        doc.styles[style].font.size = Pt(12)
        doc.styles[style].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.add_paragraph('原有风电项目初步设计', 'Title')
    doc.add_paragraph('编制单位：测试设计院')
    doc.add_paragraph('审核：测试人员')
    doc.add_paragraph('2026 年 10 月')
    doc.add_page_break()
    doc.add_paragraph('目 录')
    doc.add_paragraph('1 项目概况\t1')
    doc.add_paragraph('1.1 道路设计……2')
    doc.add_paragraph('2 主要参数\t3')
    doc.add_page_break()
    doc.add_paragraph('1 项目概况', 'Heading 1')
    doc.add_paragraph('设计标高为 10.000 m，道路长度为 20 m，净宽度为 4 m。')
    doc.add_paragraph('1.1 道路设计', 'Heading 2')
    doc.add_paragraph().add_run().add_picture(str(picture), width=Cm(13))
    doc.add_paragraph('图 1-1 道路布置示意', 'Caption')
    multi = doc.add_paragraph()
    multi.add_run('图片前的正文必须保留。')
    multi.add_run().add_picture(str(picture), width=Cm(5))
    multi.add_run('两图间文字必须保留。')
    multi.add_run().add_picture(str(picture), width=Cm(5))
    multi.add_run('图片后的正文必须保留。')
    doc.add_paragraph('2 主要参数', 'Heading 1')
    table = doc.add_table(rows=3, cols=2)
    for row, texts in zip(table.rows, [('项目', '数值'), ('标高', '10.000 m'), ('长度', '20 m')]):
        for cell, text in zip(row.cells, texts): cell.text = text
    doc.add_paragraph()
    existing = doc.add_table(rows=1, cols=1)
    existing.cell(0, 0).paragraphs[0].add_run().add_picture(str(picture), width=Cm(7))
    doc.add_paragraph('图 2-1 现有单元格图片', 'Caption')
    doc.add_paragraph('工程技术正文不应被封面或目录替换误删。')
    doc.sections[0].header.paragraphs[0].text = '工程技术正文页眉'
    doc.save(target)
    print(target)
else:
    doc = Document(target)
    with ZipFile(target) as z:
        assert not any('vbaProject' in name for name in z.namelist()), 'Macro leakage'
        xml = z.read('word/document.xml').decode()
        assert ' TOC ' in xml, 'Missing native TOC field'
        assert 'HYPERLINK' in xml or 'w:hyperlink' in xml, 'Missing TOC links'
        for text in ['图片前的正文必须保留。', '两图间文字必须保留。', '图片后的正文必须保留。', '工程技术正文不应被封面或目录替换误删。', '10.000 m', '20 m', '4 m']:
            assert text in xml, f'Missing body text: {text}'
        assert '原有风电项目初步设计' not in xml, 'Old cover not removed'
    assert len(doc.inline_shapes) == 4, f'Picture count changed: {len(doc.inline_shapes)}'
    for paragraph in doc.paragraphs:
        assert not (paragraph.style.name in ['样式1', '样式2', '样式3', '样式4'] and not paragraph.text.strip()), 'Blank numbered heading contaminates TOC'
    assert doc.paragraphs[0]._p.find('.//' + qn('w:pBdr')) is None or not any(x.get(qn('w:val')) not in ['nil', 'none'] for x in doc.paragraphs[0]._p.findall('.//' + qn('w:pBdr') + '/*')), 'Cover title border'
    image_cells = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc.findall('.//' + qn('w:drawing')):
                    image_cells.append(cell)
                    for p in cell.paragraphs:
                        spacing = p._p.find('.//' + qn('w:spacing'))
                        assert spacing is not None and spacing.get(qn('w:line')) == '240' and spacing.get(qn('w:lineRule')) == 'auto', 'Not true single spacing'
    assert len(image_cells) == 4, f'Expected each image in a cell, got {len(image_cells)}'
    assert len(doc.tables) == 5, f'Data / picture tables changed: {len(doc.tables)}'
    assert any(table.cell(1, 1).text == '10.000 m' for table in doc.tables if len(table.rows) == 3), 'Data table changed'
    print('PASS: native linked TOC, old front matter replaced, four image cells single-spaced, body and data table preserved, macro-free')
