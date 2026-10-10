"""Read-only checks for the native Word formatter's synthetic QA output."""
import sys
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document

source, output = map(Path, sys.argv[1:3])
before, after = Document(source), Document(output)
assert len(before.tables) == len(after.tables) == 1
for original, formatted in zip(before.tables, after.tables):
    assert [[c.text for c in r.cells] for r in original.rows] == [[c.text for c in r.cells] for r in formatted.rows]
for text in ['10.000 m', '20 m', '4 m', '工程格式标准化测试']:
    assert text in '\n'.join(p.text for p in after.paragraphs)
assert after.styles['样式1'] and after.styles['样式2']
with ZipFile(output) as package:
    assert not any('vbaProject' in name for name in package.namelist())
    xml = etree.fromstring(package.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    print({'paragraphs': len(after.paragraphs), 'tables': len(after.tables), 'revisions': len(xml.xpath('//w:ins|//w:del', namespaces=ns)), 'macroFree': True, 'pageWidth': after.sections[0].page_width, 'pageHeight': after.sections[0].page_height})
print('Engineering values and table contents preserved.')
