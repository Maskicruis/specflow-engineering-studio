"""Create a synthetic Word formatting fixture; never touches user documents."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

target = Path(sys.argv[1])
target.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
for style in [doc.styles['Normal'], doc.styles['Title'], doc.styles['Heading 1'], doc.styles['Heading 2']]:
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(12)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.add_paragraph('工程格式标准化测试', 'Title')
doc.add_paragraph('1 项目概况', 'Heading 1')
doc.add_paragraph('本项目采用独立文件副本进行排版测试。设计标高为 10.000 m，道路长度为 20 m。')
doc.add_paragraph('1.1 道路设计', 'Heading 2')
doc.add_paragraph('消防车道净宽度为 4 m。本段不是标题，应保留正文结构与工程数值。')
caption = doc.add_paragraph('表 1-1 主要参数', 'Caption')
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
for row, texts in zip(table.rows, [('项目', '数值'), ('设计标高', '10.000 m'), ('道路长度', '20 m')]):
    for cell, text in zip(row.cells, texts):
        cell.text = text
        cell.paragraphs[0].paragraph_format.first_line_indent = Cm(0.8)
        cell.paragraphs[0].paragraph_format.left_indent = Cm(0.4)
doc.add_paragraph('图 1-1 布置示意', 'Caption')
doc.add_paragraph('结论文本保持不变，格式修订不应编造新的设计要求。')
section.header.paragraphs[0].text = '工程设计测试'
section.header.paragraphs[0].paragraph_format.left_indent = Cm(0.8)
doc.save(target)
print(target)
